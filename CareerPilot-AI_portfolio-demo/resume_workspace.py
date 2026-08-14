import json
import re
from copy import deepcopy
from datetime import datetime
from html import escape, unescape
from io import BytesIO

import streamlit as st
import streamlit.components.v1 as components
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from lxml import html as lxml_html
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    HRFlowable,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from resume_editor_component import resume_editor

try:
    from demo_resume import PERSONAL_RESUME_BASE
except ImportError:
    PERSONAL_RESUME_BASE = None

from demo_database import (
    create_resume_version,
    get_all_experiences,
    get_experience_assets,
    get_resume_versions,
    save_resume_prompt,
    save_import_record,
)

IMPORT_START = "=== CAREERPILOT_IMPORT_START ==="
IMPORT_END = "=== CAREERPILOT_IMPORT_END ==="

PASTELS = {
    "blue": "#DCEBFF",
    "green": "#DDF3E7",
    "purple": "#EFE4FF",
    "orange": "#FFE6D2",
}

EDITOR_DEFAULTS = {
    "font": "Arial",
    "body_size": 10,
    "name_size": 21,
    "section_size": 11.5,
    "line_height": 1.18,
    "page_margin": 16,
    "section_gap": 10,
    "item_gap": 7,
    "paragraph_gap": 3,
}


def _safe(value, fallback="未设置"):
    value = fallback if value is None or not str(value).strip() else str(value).strip()
    return escape(value)


def _split_lines(value):
    if not value:
        return []
    return [line.strip(" •\t") for line in str(value).splitlines() if line.strip(" •\t")]


def _experience_payload():
    payload = []
    for exp in get_all_experiences():
        assets = get_experience_assets(exp["id"])
        payload.append({"experience": exp, "assets": assets})
    return payload


def _default_resume(payload):
    internships, projects, other = [], [], []
    for group in payload:
        exp = group["experience"]
        assets = group["assets"]
        bullets = []
        for asset in assets:
            if asset.get("include_in_resume"):
                bullets.extend(_split_lines(asset.get("resume_bullets")))
        bullets = bullets[:3] or _split_lines(exp.get("description"))[:2]
        item = {
            "source_id": exp.get("experience_code") or f'exp_{exp["id"]}',
            "organization": exp.get("organization", ""),
            "role": exp.get("role", ""),
            "date": f'{exp.get("start_date") or ""} – {exp.get("end_date") or ""}',
            "bullets": bullets,
        }
        kind = (exp.get("experience_type") or "").lower()
        if "实习" in kind or "intern" in kind:
            internships.append(item)
        elif any(key in kind for key in ("项目", "project")):
            projects.append(item)
        else:
            other.append(item)
    return {
        "name": "姓名",
        "contact": "电话｜邮箱｜LinkedIn",
        "education": [],
        "experience": internships,
        "projects": projects,
        "campus_competition_research": other,
        "skills": ["请在此补充基础简历中的真实技能"],
    }


CAPABILITY_ONTOLOGY = {
    "招聘交付": ("招聘交付", "招聘执行", "talent acquisition", "recruiting", "招聘全流程"),
    "招聘运营": ("招聘运营", "recruiting operations", "流程优化", "申请阶段", "数据汇总"),
    "候选人体验": ("候选人体验", "候选人运营", "candidate experience", "候选人沟通"),
    "人才研究": ("人才mapping", "talent mapping", "人才研究", "sourcing", "人才画像"),
    "雇主品牌": ("雇主品牌", "employer branding", "招聘传播", "内容运营", "校园招聘"),
    "HRBP与People Operations": ("hrbp", "people operations", "hr operations", "员工入职", "人力资源运营"),
    "培训与员工体验": ("培训", "training", "员工体验", "入职培训"),
    "People Analytics": ("people analytics", "数据分析", "招聘数据", "spss", "统计分析"),
    "产品策略": ("产品策略", "产品定位", "product strategy", "价值主张", "mvp"),
    "信息架构与产品设计": ("信息架构", "information architecture", "交互设计", "用户体验", "产品设计"),
    "产品开发": ("python", "streamlit", "sqlite", "产品开发", "数据库设计"),
    "HR Tech": ("hr tech", "career os", "experience bank", "人力资源技术"),
    "市场与竞争研究": ("市场研究", "竞争分析", "竞品", "市场规模", "tam", "sam", "som"),
    "商业分析与咨询": ("商业分析", "咨询", "商业模式", "策略研究", "decision support"),
    "商业表达": ("pitch deck", "路演", "商业表达", "演示设计", "英文展示"),
    "项目管理": ("项目管理", "project management", "优先级", "roadmap", "backlog", "范围管理"),
    "跨团队协作": ("跨团队", "stakeholder", "团队协作", "跨文化", "协作推进"),
    "研究设计": ("研究设计", "文献综述", "问卷", "research", "科研方法"),
    "专业写作": ("专业写作", "书稿", "案例研究", "内容重构", "编辑协作"),
    "活动运营": ("活动运营", "现场协调", "社群运营", "宣讲"),
}

CAPABILITY_WEIGHTS = {
    "招聘交付": 5, "招聘运营": 5, "候选人体验": 4, "人才研究": 4,
    "雇主品牌": 5, "HRBP与People Operations": 5, "培训与员工体验": 3,
    "People Analytics": 5, "产品策略": 5, "信息架构与产品设计": 5,
    "产品开发": 4, "HR Tech": 5, "市场与竞争研究": 5,
    "商业分析与咨询": 5, "商业表达": 4, "项目管理": 4,
    "跨团队协作": 4, "研究设计": 4, "专业写作": 3, "活动运营": 3,
}


def _detect_capabilities(text):
    normalized = str(text or "").lower()
    return {
        capability for capability, aliases in CAPABILITY_ONTOLOGY.items()
        if any(alias.lower() in normalized for alias in aliases)
    }


def _keyword_score(job, group):
    jd = " ".join(
        str(job.get(k) or "")
        for k in ("position", "details", "target_direction", "job_description")
    ).lower()
    exp = group["experience"]
    text = " ".join(str(exp.get(k) or "") for k in ("organization", "role", "skills", "summary_zh", "description")).lower()
    for asset in group["assets"]:
        text += " " + " ".join(str(asset.get(k) or "") for k in ("title", "skills", "resume_keywords", "overview_strengths")).lower()
    required = _detect_capabilities(jd)
    available = _detect_capabilities(text)
    matched = required & available

    if required:
        total_weight = sum(CAPABILITY_WEIGHTS[item] for item in required)
        matched_weight = sum(CAPABILITY_WEIGHTS[item] for item in matched)
        coverage = matched_weight / max(total_weight, 1)
    else:
        coverage = 0

    evidence_fields = sum(bool(str(asset.get(key) or "").strip()) for asset in group["assets"] for key in ("metrics", "star_result", "evidence"))
    evidence_factor = min(1.0, 0.72 + evidence_fields * 0.035)

    exp_type = str(group["experience"].get("experience_type") or "")
    auxiliary_factor = 1.0
    if "辅助" in exp_type and not required.intersection({"公众沟通与品牌", "技术叙事与科普", "跨媒介表达"}):
        auxiliary_factor = 0.22

    score = round(100 * coverage * evidence_factor * auxiliary_factor)
    matches = sorted(matched, key=lambda item: (-CAPABILITY_WEIGHTS[item], item))
    reason = f"覆盖 {len(matched)}/{len(required)} 项岗位能力" if required else "JD中尚未识别出标准能力要求"
    return score, matches[:12], reason


def _rank_experiences(job, payload):
    ranked = []
    for group in payload:
        score, matches, reason = _keyword_score(job, group)
        ranked.append({**group, "score": score, "matches": matches, "match_reason": reason})
    ranked.sort(key=lambda x: (x["score"], x["experience"].get("start_date") or ""), reverse=True)
    return ranked


def _build_prompt(job, resume, ranked, prompt_settings=None):
    prompt_settings = prompt_settings or {}
    materials = []
    for idx, group in enumerate(ranked):
        exp = group["experience"]
        level = "重点关联" if idx < 3 else "辅助关联" if idx < 6 else "其他真实素材"
        assets = []
        for asset in group["assets"]:
            assets.append({
                "source_id": asset.get("asset_code"),
                "title": asset.get("title"),
                "facts": asset.get("overview_summary") or asset.get("actions"),
                "results": asset.get("results"),
                "metrics": asset.get("metrics"),
                "resume_bullets": asset.get("resume_bullets"),
                "contribution_boundary": asset.get("contribution_boundary"),
                "expression_limits": asset.get("expression_limits"),
            })
        materials.append({
            "relevance": level,
            "source_id": exp.get("experience_code") or f'exp_{exp["id"]}',
            "organization": exp.get("organization"),
            "role": exp.get("role"),
            "date": f'{exp.get("start_date") or ""} – {exp.get("end_date") or ""}',
            "assets": assets,
        })

    schema = {
        "target_company": job.get("company", ""),
        "target_role": job.get("position", ""),
        "language": "zh-CN",
        "resume_title": f'{job.get("company", "")} {job.get("position", "")} 定制简历',
        "page_preference": "prefer_one_allow_two",
        "preserve_experience_order": True,
        "allow_remove_entire_experience": False,
        "section_order": ["education", "experience", "projects", "campus_competition_research", "skills"],
        "sections": {"education": [], "experience": [], "projects": [], "campus_competition_research": [], "skills": []},
        "omitted_content": [],
        "gaps": [],
        "items_requiring_confirmation": [],
    }

    return f"""你是一名严格遵守事实边界的中文简历优化顾问。请根据目标岗位、基础简历和 CareerPilot Experience Bank，优化一份岗位定制简历。

【目标岗位】
公司：{job.get('company', '')}
岗位：{job.get('position', '')}
截止日期：{job.get('deadline', '')}
Details：{job.get('details', '')}
完整JD：
{job.get('job_description', '')}

【当前基础简历】
{json.dumps(resume, ensure_ascii=False, indent=2)}

【本次定制偏好】
语气：{prompt_settings.get('tone', '专业、清晰')}
篇幅：{prompt_settings.get('length', '优先一页，必要时两页')}
岗位适配：根据完整 JD 自动识别关键词、职责重点、能力要求和行业表达
补充要求：{prompt_settings.get('notes', '') or '无'}

【Experience Bank真实素材，已按关联度排序】
{json.dumps(materials, ensure_ascii=False, indent=2)}

【强制规则】
1. 所有内容默认遵守职业、精炼、准确、规范四项标准：用词正式，删除重复和空泛表达，语义完整无歧义，术语、标点和格式前后一致。
2. 每条经历优先写清“做了什么、如何完成、带来什么结果”；只有素材中存在数字时才可量化，不得为了显得有成果而补造数字。
3. 根据完整 JD 自动识别关键词、职责重点、能力要求和行业表达，不要求用户重复选择通用优化方向。
4. 只能使用基础简历和 Experience Bank 中存在的信息，不得编造任何职责、成果、数字、工具、证书或技能。
5. 公司名称、学校名称、岗位名称、项目名称和起止时间必须严格沿用，不得修改。
6. 不得把团队成果写成个人独立完成，不得改变责任边界。
7. 不得删除任何整段已有经历；弱相关经历可以压缩，但至少保留1条bullet。
8. 所有经历保持由近到远，不得调整时间顺序。
9. 优先控制一页，合理压缩后仍放不下时允许两页。
10. 数字数值不得改变；不确定内容放入 items_requiring_confirmation。
11. 当前只生成中文简历。
12. 基础简历的版式与结构完全锁定：不得新增“个人概述”、求职方向或任何新章节；不得改变姓名与基础信息的单行排列；不得改变教育经历的学校、学院/专业、地点、日期和补充说明所在行；不得新增、删除或重排经历条目和 bullet。只允许为原有 bullet 提供逐条对应的优化文本。
13. 横线、章节顺序、条目顺序、每条经历的行结构、字体、字号、行距与页边距均由 CareerPilot 原模板控制，输出内容不得建议或携带新的排版。

【输出顺序】
1. JD核心要求
2. 修改策略与主要变化
3. 完整中文定制简历
4. 严格输出以下导入区块，且每段经历必须保留 source_id：

{IMPORT_START}
{json.dumps(schema, ensure_ascii=False, indent=2)}
{IMPORT_END}
"""


def _plain_ai_resume_html(text):
    """Turn a plain Markdown-ish AI resume into safe, editable A4 HTML."""
    cleaned = re.sub(r"^```(?:markdown|md|text)?\s*|\s*```$", "", (text or "").strip(), flags=re.IGNORECASE)
    lines = cleaned.splitlines()
    resume_markers = ("完整中文定制简历", "完整定制简历", "定制后简历", "优化后简历")
    for index, line in enumerate(lines):
        normalized = re.sub(r"[#*：:\s]", "", line)
        if any(marker in normalized for marker in resume_markers):
            lines = lines[index + 1:]
            break

    def inline(value):
        safe = _safe(value, "")
        return re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", safe)

    section_names = (
        "教育经历", "教育背景", "实习经历", "工作经历", "项目经历", "校园经历",
        "竞赛经历", "科研经历", "校园 / 竞赛 / 科研", "技能", "专业技能",
        "语言能力", "证书", "荣誉奖项", "个人总结",
    )
    html_parts = []
    list_open = False

    def close_list():
        nonlocal list_open
        if list_open:
            html_parts.append("</ul>")
            list_open = False

    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            close_list()
            continue
        if line in {IMPORT_START, IMPORT_END}:
            break
        line = re.sub(r"^>\s?", "", line)
        heading = re.match(r"^(#{1,6})\s*(.+)$", line)
        if heading:
            close_list()
            level, title = len(heading.group(1)), heading.group(2).strip().strip("* ")
            html_parts.append(f"<h1>{inline(title)}</h1>" if level == 1 else f"<h2>{inline(title)}</h2>")
            continue
        plain_heading = line.strip("* ").rstrip("：:")
        if plain_heading in section_names:
            close_list()
            html_parts.append(f"<h2>{inline(plain_heading)}</h2>")
            continue
        bullet = re.match(r"^(?:[-*•]|\d+[.)])\s+(.+)$", line)
        if bullet:
            if not list_open:
                html_parts.append("<ul>")
                list_open = True
            html_parts.append(f"<li>{inline(bullet.group(1))}</li>")
            continue
        close_list()
        html_parts.append(f"<p>{inline(line)}</p>")
    close_list()
    if not html_parts:
        raise ValueError("AI 结果为空，暂时没有可导入的简历内容。")
    return "".join(html_parts)


def _parse_import(text):
    if not (text or "").strip():
        raise ValueError("请先粘贴 AI 生成的简历结果。")
    if IMPORT_START not in text or IMPORT_END not in text:
        return {
            "sections": {},
            "items_requiring_confirmation": [],
            "_import_mode": "plain_text",
            "_plain_text": text,
            "_plain_html": _plain_ai_resume_html(text),
        }
    raw = text.split(IMPORT_START, 1)[1].split(IMPORT_END, 1)[0].strip()
    data = json.loads(raw)
    if not isinstance(data.get("sections"), dict):
        raise ValueError("JSON 缺少 sections。")
    for section in ("experience", "projects", "campus_competition_research"):
        for item in data["sections"].get(section, []):
            if not item.get("source_id"):
                raise ValueError(f"{section} 中存在缺少 source_id 的经历。")
    return data


def _resume_from_import(data, fallback):
    if fallback.get("sections"):
        result = deepcopy(fallback)

        def normalized(value):
            value = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", str(value or ""))
            return re.sub(r"[^0-9a-zA-Z\u4e00-\u9fff]+", "", value).lower()

        def clean_ai_text(value):
            value = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", str(value or ""))
            value = re.sub(r"[*_`#]", "", value)
            return re.sub(r"\s+", " ", value).strip(" -•\t")

        def bullet_value(value):
            if isinstance(value, dict):
                return value.get("tailored") or value.get("text") or value.get("original") or ""
            return str(value or "")

        imported_entries = []
        for section_entries in (data.get("sections") or {}).values():
            if isinstance(section_entries, list):
                imported_entries.extend(item for item in section_entries if isinstance(item, dict))

        plain_text = str(data.get("_plain_text") or "")
        all_boundaries = []
        for section in result.get("sections", []):
            all_boundaries.append(section.get("title") or "")
            all_boundaries.extend(entry.get("title") or "" for entry in section.get("entries", []))

        def plain_bullets_for(title):
            if not plain_text or not title:
                return []
            variants = [title]
            if "｜" in title:
                variants.append(title.split("｜", 1)[0])
            start = -1
            matched = ""
            for variant in variants:
                start = plain_text.find(variant)
                if start >= 0:
                    matched = variant
                    break
            if start < 0:
                return []
            block_start = start + len(matched)
            end = len(plain_text)
            for boundary in all_boundaries:
                if not boundary or boundary == title:
                    continue
                boundary_variants = [boundary, boundary.split("｜", 1)[0]]
                for candidate in boundary_variants:
                    position = plain_text.find(candidate, block_start)
                    if position >= 0:
                        end = min(end, position)
            block = plain_text[block_start:end]
            collected = []
            current = ""
            for raw_line in block.splitlines():
                line = raw_line.strip()
                bullet = re.match(r"^(?:[-*•]|\d+[.)])\s+(.+)$", line)
                if bullet:
                    if current:
                        collected.append(clean_ai_text(current))
                    current = bullet.group(1)
                elif current and line and not re.match(r"^#{1,6}\s", line):
                    current += " " + line
            if current:
                collected.append(clean_ai_text(current))
            return [item for item in collected if item]

        for section in result.get("sections", []):
            # Education and profile are factual layout anchors and never change on AI import.
            if section.get("id") == "education":
                continue
            for entry in section.get("entries", []):
                title_key = normalized(entry.get("title"))
                replacement_bullets = []
                for imported in imported_entries:
                    imported_title = (
                        imported.get("title") or imported.get("organization") or imported.get("company")
                        or imported.get("project") or imported.get("school") or ""
                    )
                    imported_key = normalized(imported_title)
                    if title_key and imported_key and (title_key == imported_key or title_key.startswith(imported_key) or imported_key.startswith(title_key)):
                        replacement_bullets = [clean_ai_text(bullet_value(item)) for item in imported.get("bullets", [])]
                        break
                if not replacement_bullets:
                    replacement_bullets = plain_bullets_for(entry.get("title") or "")
                if not replacement_bullets:
                    continue
                original_bullets = entry.get("bullets") or []
                merged = []
                for index, original in enumerate(original_bullets):
                    if index >= len(replacement_bullets):
                        merged.append(original)
                        continue
                    if isinstance(original, dict):
                        updated = deepcopy(original)
                        updated["tailored"] = replacement_bullets[index]
                        merged.append(updated)
                    else:
                        merged.append(replacement_bullets[index])
                entry["bullets"] = merged
        return result

    sections = data.get("sections", {})
    fallback_profile = fallback.get("profile", {})
    return {
        "name": fallback.get("name") or fallback_profile.get("name") or "姓名",
        "contact": fallback.get("contact") or fallback_profile.get("contact") or "电话｜邮箱",
        "education": sections.get("education", fallback.get("education", [])),
        "experience": sections.get("experience", fallback.get("experience", [])),
        "projects": sections.get("projects", fallback.get("projects", [])),
        "campus_competition_research": sections.get("campus_competition_research", fallback.get("campus_competition_research", [])),
        "skills": sections.get("skills", fallback.get("skills", [])),
    }


def _legacy_editor_import_data(editor_html):
    """Recover entry bullets from an older free-form imported editor document."""
    try:
        root = lxml_html.fragment_fromstring(editor_html or "", create_parent=True)
    except Exception:
        return {"sections": {"legacy": []}}
    entries = []
    for heading in root.xpath(".//h2"):
        title = re.sub(r"\s+", " ", heading.text_content() or "").strip()
        bullets = []
        sibling = heading.getnext()
        while sibling is not None and str(sibling.tag).lower() != "h2":
            if str(sibling.tag).lower() == "ul":
                bullets.extend(
                    re.sub(r"\s+", " ", item.text_content() or "").strip()
                    for item in sibling.xpath(".//li")
                )
            sibling = sibling.getnext()
        if title and bullets:
            entries.append({"title": title, "bullets": [item for item in bullets if item]})
    return {"sections": {"legacy": entries}}


def _resume_html(resume):
    def bullet_text(value):
        if isinstance(value, dict):
            return value.get("tailored") or value.get("original") or value.get("text") or ""
        return str(value or "")

    def items_html(items, org_key="organization"):
        rows = []
        for item in items or []:
            if isinstance(item, str):
                rows.append(f'<p>{_safe(item)}</p>')
                continue
            title = item.get(org_key) or item.get("company") or item.get("project") or item.get("school") or ""
            role = item.get("subtitle") or item.get("role") or item.get("position") or item.get("degree") or ""
            date = item.get("date") or item.get("time") or ""
            location = item.get("location") or ""
            role_html = _safe(role, "")
            subtitle_url = str(item.get("subtitle_url") or "").strip()
            if role and re.match(r"^https?://", subtitle_url, flags=re.IGNORECASE):
                role_html = (
                    f'<a class="inline-link" href="{_safe(subtitle_url)}" '
                    f'target="_blank" rel="noopener noreferrer">{_safe(role)}</a>'
                )
            detail_values = item.get("details", [])
            if isinstance(detail_values, str):
                detail_values = [detail_values]
            details = "".join(f"<p>{_safe(d)}</p>" for d in detail_values if d)
            link_url = str(item.get("link") or "").strip()
            link_html = ""
            if re.match(r"^https?://", link_url, flags=re.IGNORECASE):
                link_label = item.get("link_label") or link_url
                link_prefix = item.get("link_prefix") or "Demo"
                link_html = (
                    f'<p class="entry-link">{_safe(link_prefix)}：<a href="{_safe(link_url)}" '
                    f'target="_blank" rel="noopener noreferrer">{_safe(link_label)}</a></p>'
                )
            bullets = "".join(f"<li>{_safe(bullet_text(b))}</li>" for b in item.get("bullets", []) if bullet_text(b))
            row_right = location or (date if not role else "")
            subrow_right = date if (role or location) else ""
            subrow = ""
            if role or subrow_right:
                subrow = (
                    f'<div class="subrow"><span class="left">{role_html}</span>'
                    f'<span class="right">{_safe(subrow_right, "")}</span></div>'
                )
            rows.append(
                f'<div class="entry"><div class="row"><strong>{_safe(title)}</strong>'
                f'<span class="right">{_safe(row_right, "")}</span></div>{subrow}{details}{link_html}'
                f'<ul>{bullets}</ul></div>'
            )
        return "".join(rows) or '<p style="color:#999">该模块暂未录入</p>'

    if resume.get("sections"):
        profile = resume.get("profile", {})
        sections = []
        for section in resume.get("sections", []):
            sections.append(
                f'<h2>{_safe(section.get("title"), "未命名模块")}</h2>'
                f'{items_html(section.get("entries", []), "title")}'
            )
        return (
            f'<h1>{_safe(profile.get("name"), "姓名")}</h1>'
            f'<p class="contact">{_safe(profile.get("meta"), "")}</p>'
            f'<p class="contact contact-secondary">{_safe(profile.get("contact"), "")}</p>'
            + "".join(sections)
        )

    skills = resume.get("skills", [])
    if isinstance(skills, dict):
        skills = [f"{k}：{v}" for k, v in skills.items()]
    return f'''<h1>{_safe(resume.get('name'), '姓名')}</h1>
      <p class="contact">{_safe(resume.get('contact'), '电话｜邮箱')}</p>
      <h2>教育经历</h2>{items_html(resume.get('education', []), 'school')}
      <h2>实习经历</h2>{items_html(resume.get('experience', []))}
      <h2>项目经历</h2>{items_html(resume.get('projects', []), 'project')}
      <h2>校园 / 竞赛 / 科研</h2>{items_html(resume.get('campus_competition_research', []), 'organization')}
      <h2>技能</h2><p class="skills">{' · '.join(_safe(x) for x in skills)}</p>'''


def _editor_value(resume):
    return resume.get("_editor_html") or _resume_html(resume)


def _editor_settings(resume):
    return {**EDITOR_DEFAULTS, **(resume.get("_editor_settings") or {})}


def _copy_prompt_button(prompt):
    prompt_json = json.dumps(prompt, ensure_ascii=False).replace("</", "<\\/")
    components.html(
        f"""
        <button id="copy-prompt" type="button">复制 Prompt</button>
        <script>
        const promptText = {prompt_json};
        const button = document.getElementById('copy-prompt');
        button.addEventListener('click', async () => {{
            let copied = false;
            try {{
                await navigator.clipboard.writeText(promptText);
                copied = true;
            }} catch (error) {{
                const field = document.createElement('textarea');
                field.value = promptText;
                field.style.position = 'fixed';
                field.style.opacity = '0';
                document.body.appendChild(field);
                field.focus();
                field.select();
                copied = document.execCommand('copy');
                field.remove();
            }}
            button.textContent = copied ? '已复制到剪贴板' : '复制失败，请手动选择';
            button.classList.toggle('copied', copied);
            window.setTimeout(() => {{
                button.textContent = '复制 Prompt';
                button.classList.remove('copied');
            }}, 1800);
        }});
        </script>
        <style>
        html,body{{margin:0;padding:0;background:transparent;font-family:"Noto Sans SC","PingFang SC",sans-serif}}
        #copy-prompt{{width:100%;height:38px;border:1px solid #d98249;border-radius:10px;background:#dd7f46;color:#fff;font-size:13px;font-weight:750;cursor:pointer;transition:background .16s ease,transform .16s ease}}
        #copy-prompt:hover{{background:#c96d38;transform:translateY(-1px)}}
        #copy-prompt.copied{{background:#efb07d;border-color:#efb07d;color:#603719}}
        </style>
        """,
        height=42,
    )


def _apply_editor_result(resume, result):
    updated = deepcopy(resume)
    if isinstance(result, str):
        updated["_editor_html"] = result
    elif isinstance(result, dict):
        updated["_editor_html"] = result.get("html") or _editor_value(resume)
        updated["_editor_settings"] = {**EDITOR_DEFAULTS, **(result.get("settings") or {})}
        updated["_editor_page_count"] = max(1, int(result.get("page_count") or 1))
    return updated


def _text_units(value):
    text = re.sub(r"<[^>]+>", "\n", value or "")
    return [
        re.sub(r"\s+", " ", unescape(line)).strip()
        for line in text.splitlines()
        if re.sub(r"\s+", " ", unescape(line)).strip()
    ]


def _change_summary(base_html, generated_html):
    base_units = _text_units(base_html)
    generated_units = _text_units(generated_html)
    base_set, generated_set = set(base_units), set(generated_units)
    return {
        "added": [item for item in generated_units if item not in base_set][:18],
        "removed": [item for item in base_units if item not in generated_set][:18],
    }


def _prompt_notes(general_note, comments):
    """Merge free-form guidance and review comments into one prompt section."""
    parts = []
    if str(general_note or "").strip():
        parts.append(f"整体备注：{str(general_note).strip()}")
    for index, comment in enumerate(comments or [], start=1):
        quote = str(comment.get("quote") or "").strip()
        note = str(comment.get("note") or "").strip()
        if quote and note:
            parts.append(f"批注 {index}\n原文：{quote}\n修改要求：{note}")
    return "\n\n".join(parts)


def _set_run_font(run, font_name, size, bold=False):
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def _set_paragraph_border(paragraph, color="333333", size="6"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def _remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    tbl_pr.append(borders)


def _docx_bytes(resume):
    """Create a real A4 DOCX from the clean editor HTML."""
    settings = _editor_settings(resume)
    font_name = settings.get("font") or "Arial"
    body_size = float(settings.get("body_size") or 10)
    paragraph_gap = float(settings.get("paragraph_gap") or 3) * 0.75
    line_height = float(settings.get("line_height") or 1.18)
    page_margin = float(settings.get("page_margin") or 16)
    root = lxml_html.fragment_fromstring(_editor_value(resume), create_parent="div")

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(page_margin)
    section.left_margin = section.right_margin = Mm(page_margin)
    normal = document.styles["Normal"]
    normal.font.name = font_name
    normal.font.size = Pt(body_size)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)

    def format_paragraph(paragraph, after=paragraph_gap):
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(after)
        paragraph.paragraph_format.line_spacing = line_height

    def add_text_paragraph(text, bold=False, align=None, size=body_size, style=None):
        paragraph = document.add_paragraph(style=style)
        if align is not None:
            paragraph.alignment = align
        format_paragraph(paragraph)
        run = paragraph.add_run(re.sub(r"\s+", " ", text or "").strip())
        _set_run_font(run, font_name, size, bold)
        return paragraph

    def add_entry(entry):
        row_node = entry.xpath('./div[contains(concat(" ", normalize-space(@class), " "), " row ")]')
        subrow_node = entry.xpath('./div[contains(concat(" ", normalize-space(@class), " "), " subrow ")]')
        for idx, nodes in enumerate((row_node, subrow_node)):
            if not nodes:
                continue
            node = nodes[0]
            children = list(node)
            parts = [
                re.sub(r"\s+", " ", children[index].text_content()).strip()
                if index < len(children) else ""
                for index in range(2)
            ]
            table = document.add_table(rows=1, cols=2)
            table.autofit = False
            _remove_table_borders(table)
            left_text = parts[0] if parts else ""
            right_text = parts[-1] if len(parts) > 1 else ""
            for col, text in enumerate((left_text, right_text)):
                paragraph = table.cell(0, col).paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if col else WD_ALIGN_PARAGRAPH.LEFT
                format_paragraph(paragraph, 0)
                run = paragraph.add_run(text)
                _set_run_font(run, font_name, body_size, bold=(idx == 0))
        for paragraph_node in entry.xpath('./p'):
            add_text_paragraph(paragraph_node.text_content())
        for item in entry.xpath('./ul/li|./ol/li'):
            add_text_paragraph(item.text_content(), style="List Bullet")

    for node in root:
        tag = (node.tag or "").lower() if isinstance(node.tag, str) else ""
        classes = set((node.get("class") or "").split())
        text = re.sub(r"\s+", " ", node.text_content()).strip()
        if tag == "h1":
            add_text_paragraph(text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=float(settings.get("name_size") or 21))
        elif tag == "h2":
            paragraph = add_text_paragraph(text, bold=True, size=float(settings.get("section_size") or 11.5))
            paragraph.paragraph_format.space_before = Pt(float(settings.get("section_gap") or 10) * 0.75)
            _set_paragraph_border(paragraph)
        elif tag == "p":
            align = WD_ALIGN_PARAGRAPH.CENTER if "contact" in classes else WD_ALIGN_PARAGRAPH.LEFT
            add_text_paragraph(text, align=align)
        elif tag == "div" and "entry" in classes:
            add_entry(node)
            document.paragraphs[-1].paragraph_format.space_after = Pt(float(settings.get("item_gap") or 7) * 0.75) if document.paragraphs else Pt(0)
        elif tag in {"ul", "ol"}:
            for item in node.xpath('./li'):
                add_text_paragraph(item.text_content(), style="List Bullet")
        elif tag == "div" and "page-break" in classes:
            document.add_page_break()

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf_bytes(resume):
    """Create a clean, downloadable A4 PDF from the editor HTML."""
    settings = _editor_settings(resume)
    body_size = float(settings.get("body_size") or 10)
    name_size = float(settings.get("name_size") or 21)
    section_size = float(settings.get("section_size") or 11.5)
    line_height = float(settings.get("line_height") or 1.18)
    page_margin = float(settings.get("page_margin") or 16)
    paragraph_gap = float(settings.get("paragraph_gap") or 3)
    section_gap = float(settings.get("section_gap") or 10)
    item_gap = float(settings.get("item_gap") or 7)
    root = lxml_html.fragment_fromstring(_editor_value(resume), create_parent="div")

    font_name = "CareerPilotPDF"
    font_candidates = (
        "/System/Library/Fonts/Songti.ttc",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/System/Library/Fonts/STHeiti Medium.ttc",
    )
    if font_name not in pdfmetrics.getRegisteredFontNames():
        registered = False
        for font_path in font_candidates:
            try:
                pdfmetrics.registerFont(TTFont(font_name, font_path, subfontIndex=0))
                pdfmetrics.registerFontFamily(
                    font_name,
                    normal=font_name,
                    bold=font_name,
                    italic=font_name,
                    boldItalic=font_name,
                )
                registered = True
                break
            except Exception:
                continue
        if not registered:
            font_name = "STSong-Light"
            if font_name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(UnicodeCIDFont(font_name))

    body_style = ParagraphStyle(
        "ResumeBody",
        fontName=font_name,
        fontSize=body_size,
        leading=body_size * line_height,
        textColor=colors.HexColor("#161616"),
        alignment=TA_LEFT,
        spaceBefore=0,
        spaceAfter=paragraph_gap,
        allowWidows=0,
        allowOrphans=0,
    )
    name_style = ParagraphStyle(
        "ResumeName",
        parent=body_style,
        fontSize=name_size,
        leading=name_size * 1.08,
        alignment=TA_CENTER,
        spaceAfter=5,
    )
    contact_style = ParagraphStyle(
        "ResumeContact",
        parent=body_style,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#333333"),
        spaceAfter=8,
    )
    section_style = ParagraphStyle(
        "ResumeSection",
        parent=body_style,
        fontSize=section_size,
        leading=section_size * 1.2,
        spaceBefore=section_gap,
        spaceAfter=2,
    )
    row_left_style = ParagraphStyle("ResumeRowLeft", parent=body_style, spaceAfter=0)
    row_right_style = ParagraphStyle("ResumeRowRight", parent=body_style, alignment=TA_RIGHT, spaceAfter=0)
    bullet_style = ParagraphStyle(
        "ResumeBullet",
        parent=body_style,
        leftIndent=12,
        firstLineIndent=-8,
        spaceAfter=paragraph_gap,
    )

    def clean_text(value):
        return escape(re.sub(r"\s+", " ", value or "").strip())

    def text_paragraph(value, style=body_style, bold=False):
        content = clean_text(value)
        return Paragraph(f"<b>{content}</b>" if bold else content, style)

    def entry_flowables(entry):
        parts = []
        row_node = entry.xpath('./div[contains(concat(" ", normalize-space(@class), " "), " row ")]')
        subrow_node = entry.xpath('./div[contains(concat(" ", normalize-space(@class), " "), " subrow ")]')
        for index, nodes in enumerate((row_node, subrow_node)):
            if not nodes:
                continue
            node = nodes[0]
            children = list(node)
            left_text = re.sub(r"\s+", " ", children[0].text_content()).strip() if children else ""
            right_text = re.sub(r"\s+", " ", children[1].text_content()).strip() if len(children) > 1 else ""
            left = text_paragraph(left_text, row_left_style, bold=(index == 0))
            right = text_paragraph(right_text, row_right_style, bold=(index == 0))
            table = Table([[left, right]], colWidths=[None, 44 * mm], hAlign="LEFT")
            table.setStyle(TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
            ]))
            parts.append(table)
        for paragraph_node in entry.xpath('./p'):
            parts.append(text_paragraph(paragraph_node.text_content()))
        for item in entry.xpath('./ul/li|./ol/li'):
            parts.append(Paragraph(f"• {clean_text(item.text_content())}", bullet_style))
        parts.append(Spacer(1, item_gap))
        return parts

    story = []
    pending_section = []
    content_width = A4[0] - (2 * page_margin * mm)

    def section_flowable(value):
        block = Table(
            [[text_paragraph(value, section_style, bold=True)]],
            colWidths=[content_width],
            hAlign="LEFT",
        )
        block.setStyle(TableStyle([
            ("LINEBELOW", (0, 0), (-1, -1), 0.8, colors.HexColor("#333333")),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 0),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        block.keepWithNext = 1
        return block

    for node in root:
        tag = (node.tag or "").lower() if isinstance(node.tag, str) else ""
        classes = set((node.get("class") or "").split())
        text = re.sub(r"\s+", " ", node.text_content()).strip()
        if not text and not (tag == "div" and "page-break" in classes):
            continue
        if tag == "h1":
            story.append(text_paragraph(text, name_style, bold=True))
        elif tag == "h2":
            if pending_section:
                story.extend(pending_section)
            pending_section = [section_flowable(text)]
        elif tag == "p":
            paragraph = text_paragraph(text, contact_style if "contact" in classes else body_style)
            if pending_section:
                story.append(KeepTogether(pending_section + [paragraph]))
                pending_section = []
            else:
                story.append(paragraph)
        elif tag == "div" and "entry" in classes:
            flows = entry_flowables(node)
            story.append(KeepTogether(pending_section + flows))
            pending_section = []
        elif tag in {"ul", "ol"}:
            flows = [
                Paragraph(f"• {clean_text(item.text_content())}", bullet_style)
                for item in node.xpath('./li')
            ]
            if pending_section:
                story.append(KeepTogether(pending_section + flows))
                pending_section = []
            else:
                story.extend(flows)
        elif tag == "div" and "page-break" in classes:
            if pending_section:
                story.extend(pending_section)
                pending_section = []
            story.append(PageBreak())

    if pending_section:
        story.extend(pending_section)

    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=page_margin * mm,
        leftMargin=page_margin * mm,
        topMargin=page_margin * mm,
        bottomMargin=page_margin * mm,
        title="CareerPilot 定制简历",
        author="CareerPilot",
    )
    document.build(story)
    return output.getvalue()


def _interview_pack(job, resume, ranked):
    questions = []
    stories = []
    for group in ranked[:6]:
        exp = group["experience"]
        for asset in group["assets"]:
            if not asset.get("include_in_interview"):
                continue
            for q in _split_lines(asset.get("interview_questions"))[:2]:
                questions.append({"question": q, "source": asset.get("asset_code"), "experience": exp.get("organization")})
            if asset.get("star_situation"):
                stories.append({
                    "source": asset.get("asset_code"),
                    "title": asset.get("title"),
                    "experience": exp.get("organization"),
                    "s": asset.get("star_situation"), "t": asset.get("star_task"),
                    "a": asset.get("star_action"), "r": asset.get("star_result"),
                    "followups": _split_lines(asset.get("interview_followups"))[:3],
                    "risk": asset.get("interview_risks") or asset.get("expression_limits"),
                })
    return {"job": job, "questions": questions[:18], "stories": stories[:10]}


def _legacy_styles():
    st.markdown("""
    <style>
    .rw-heading{display:flex;align-items:flex-end;justify-content:space-between;gap:20px;margin:2px 0 14px}
    .rw-heading h1{font-size:28px;line-height:1.2;letter-spacing:-.04em;margin:0;color:#171717}.rw-heading p{margin:5px 0 0;color:#737373;font-size:13px}
    .rw-kicker{font-size:10px;letter-spacing:.13em;text-transform:uppercase;color:#8a8a8a;font-weight:800;margin-bottom:7px}
    .rw-info-grid{display:grid;grid-template-columns:1.5fr repeat(4,minmax(105px,1fr));gap:9px;margin:12px 0 18px}
    .rw-info{min-height:68px;border:1px solid #e1e1df;border-radius:13px;background:#fafaf9;padding:11px 13px;display:flex;flex-direction:column;justify-content:center}
    .rw-info small{display:block;color:#858585;font-size:10px;margin-bottom:5px}.rw-info b{font-size:12px;line-height:1.35;color:#202020}.rw-info:first-child{background:#1b1b1b;border-color:#1b1b1b}.rw-info:first-child small,.rw-info:first-child b{color:#fff}
    .rw-editor-title{display:flex;align-items:center;justify-content:space-between;margin:0 0 8px;padding:0 2px}.rw-editor-title b{font-size:14px}.rw-editor-title span{font-size:10px;color:#7a7a7a;border:1px solid #dedede;border-radius:99px;padding:4px 8px;background:#fafafa}
    .rw-section-head{display:flex;align-items:end;justify-content:space-between;margin:27px 0 12px;border-top:1px solid #dedede;padding-top:20px}.rw-section-head h2{font-size:20px;letter-spacing:-.025em;margin:0}.rw-section-head p{font-size:11px;color:#777;margin:0}
    .rw-material-track{display:flex;gap:10px;overflow-x:auto;padding:2px 2px 13px;scroll-snap-type:x proximity;scrollbar-width:thin;scrollbar-color:#bcbcbc #eeeeec}
    .rw-material{flex:0 0 245px;min-height:124px;scroll-snap-align:start;border:1px solid #dededb;border-radius:15px;background:#fafaf9;padding:14px;display:flex;flex-direction:column}
    .rw-material-top{display:flex;justify-content:space-between;gap:10px;margin-bottom:12px}.rw-material-code{font-size:9px;color:#777;letter-spacing:.06em}.rw-score{font-size:9px;color:#333;background:#e9e9e6;border-radius:99px;padding:3px 7px}.rw-material b{font-size:13px;line-height:1.35}.rw-material p{font-size:11px;color:#707070;line-height:1.45;margin:5px 0 10px}.rw-tags{margin-top:auto;font-size:9px;color:#555}
    .rw-panel-label{font-size:11px;text-transform:uppercase;letter-spacing:.1em;font-weight:800;color:#707070;margin:2px 0 8px}.rw-panel-copy{font-size:12px;line-height:1.6;color:#696969;margin:0 0 14px}
    .rw-summary{padding:11px 13px;border:1px solid #dededb;border-radius:12px;background:#f6f6f4;font-size:11px;line-height:1.65;margin:8px 0}
    .rw-interview{margin-top:18px;border:1px solid #dededb;border-radius:18px;padding:18px;background:#f7f7f5}.rw-q{background:#fff;border:1px solid #e6e6e3;padding:12px;border-radius:11px;margin:8px 0;font-size:12px}.rw-star{background:#fff;border:1px solid #e6e6e3;padding:15px;border-radius:12px;margin:10px 0}.rw-star b{color:#222}.rw-star p{font-size:11px;line-height:1.65;margin:5px 0}
    div[data-testid="stCustomComponentV1"]{border-radius:18px;overflow:hidden}div[data-testid="stVerticalBlockBorderWrapper"]{border-color:#dededb!important;border-radius:16px!important;background:#fafaf9!important}
    @media(max-width:1100px){.rw-info-grid{grid-template-columns:repeat(2,1fr)}.rw-info:first-child{grid-column:1/-1}.rw-heading{align-items:flex-start;flex-direction:column}}
    </style>
    """, unsafe_allow_html=True)


def _render_resume_workspace_legacy(jobs):
    _legacy_styles()
    payload = _experience_payload()
    if not jobs:
        st.warning("请先在“职位申请”中创建目标岗位。")
        return

    st.markdown('''<div class="rw-heading"><div><div class="rw-kicker">Resume tailoring studio</div><h1>简历定制</h1><p>选择岗位，对照编辑基础简历与岗位版本，再用真实经历素材完成定制。</p></div></div>''', unsafe_allow_html=True)

    job_labels = {f'{j["company"]}｜{j["position"]}': j for j in jobs}
    selected_label = st.selectbox("选择目标岗位 / JD", list(job_labels), key="rw_job")
    job = job_labels[selected_label]
    ranked = _rank_experiences(job, payload)

    versions = get_resume_versions(job["id"])
    base_version = next((v for v in versions if v.get("version_type") == "base_edit"), None)
    generated_version = next((v for v in versions if v.get("version_type") != "base_edit"), None)

    base_key = f'rw_base_{job["id"]}'
    generated_key = f'rw_generated_{job["id"]}'
    editor_revision_key = f"rw_editor_revision_{job['id']}"
    if editor_revision_key not in st.session_state:
        st.session_state[editor_revision_key] = 0
    if base_key not in st.session_state:
        if base_version:
            st.session_state[base_key] = json.loads(base_version["content_json"])
        elif PERSONAL_RESUME_BASE:
            st.session_state[base_key] = deepcopy(PERSONAL_RESUME_BASE)
        else:
            st.session_state[base_key] = _default_resume(payload)
    if generated_key not in st.session_state:
        source = json.loads(generated_version["content_json"]) if generated_version else deepcopy(st.session_state[base_key])
        st.session_state[generated_key] = source
    base_resume = st.session_state[base_key]
    generated_resume = st.session_state[generated_key]
    template_resume = base_resume if base_resume.get("sections") else (deepcopy(PERSONAL_RESUME_BASE) if PERSONAL_RESUME_BASE else None)
    current_editor_html = _editor_value(generated_resume)
    needs_template_repair = (
        not generated_resume.get("sections")
        or "个人概述" in current_editor_html
        or "mailto:" in current_editor_html
    )
    if template_resume and needs_template_repair:
        migrated = _resume_from_import(
            _legacy_editor_import_data(current_editor_html),
            template_resume,
        )
        migrated["_editor_settings"] = _editor_settings(generated_resume)
        migrated["_review_comments"] = deepcopy(generated_resume.get("_review_comments") or [])
        migrated["_editor_html"] = _resume_html(migrated)
        st.session_state[generated_key] = migrated
        generated_resume = migrated
        st.session_state[editor_revision_key] += 1

    current_version_name = generated_version.get("version_name") if generated_version else "待生成"
    base_version_name = base_version.get("version_name") if base_version else "个人基础简历"
    text_length = len(re.sub(r"<[^>]+>", "", _editor_value(generated_resume)))
    page_estimate = "约 1 页" if text_length < 5000 else "约 2 页"

    st.markdown(f'''
    <div class="rw-info-grid">
      <div class="rw-info"><small>目标岗位</small><b>{_safe(job.get('company'))}｜{_safe(job.get('position'))}</b></div>
      <div class="rw-info"><small>截止日期</small><b>{_safe(job.get('deadline'))}</b></div>
      <div class="rw-info"><small>基础简历</small><b>{_safe(base_version_name)}</b></div>
      <div class="rw-info"><small>当前版本</small><b>{_safe(current_version_name)}</b></div>
      <div class="rw-info"><small>预计页数</small><b>{page_estimate}</b></div>
    </div>
    ''', unsafe_allow_html=True)

    left, right = st.columns(2, gap="medium")
    with left:
        st.markdown('<div class="rw-editor-title"><b>基础简历</b><span>原始内容 · 可直接编辑</span></div>', unsafe_allow_html=True)
        base_result = resume_editor(
            _editor_value(base_resume), _editor_settings(base_resume),
            key=f"rw_base_editor_{job['id']}", height=820,
        )
        st.session_state[base_key] = _apply_editor_result(base_resume, base_result)
        if st.button("保存基础简历修改", key=f"rw_save_base_{job['id']}", use_container_width=True):
            create_resume_version(job["id"], f"基础简历 {datetime.now():%m%d %H:%M}", st.session_state[base_key], "base_edit")
            st.success("基础简历及排版设置已保存。")

    with right:
        st.markdown('<div class="rw-editor-title"><b>生成后简历</b><span>岗位版本 · 可继续精修</span></div>', unsafe_allow_html=True)
        generated_result = resume_editor(
            _editor_value(generated_resume), _editor_settings(generated_resume),
            key=f"rw_generated_editor_{job['id']}", height=820,
        )
        st.session_state[generated_key] = _apply_editor_result(generated_resume, generated_result)
        if st.button("保存当前岗位版本", key=f"rw_save_generated_{job['id']}", type="primary", use_container_width=True):
            create_resume_version(job["id"], f'{job["company"]} {job["position"]} 手动版 {datetime.now():%m%d %H:%M}', st.session_state[generated_key], "manual")
            st.success("岗位版本及排版设置已保存。")

    st.markdown('<div class="rw-section-head"><div><div class="rw-kicker">Source material</div><h2>Experience Bank 素材</h2></div><p>左右滑动浏览 · 按 JD 相关度排序</p></div>', unsafe_allow_html=True)
    material_cards = []
    experience_options = {}
    for group in ranked:
        exp = group["experience"]
        code = exp.get("experience_code") or f'exp_{exp["id"]}'
        label = f'{exp.get("organization") or "未命名经历"}｜{exp.get("role") or "经历"} [{code}]'
        experience_options[label] = group
        tags = " · ".join(group.get("matches") or []) or "等待 JD 关键词匹配"
        material_cards.append(f'''<div class="rw-material"><div class="rw-material-top"><span class="rw-material-code">{_safe(code)}</span><span class="rw-score">匹配度 {group.get("score", 0)}%</span></div><b>{_safe(exp.get("organization"))}</b><p>{_safe(exp.get("role"))}</p><div class="rw-tags">{_safe(group.get("match_reason"))}<br>{_safe(tags)}</div></div>''')
    st.markdown('<div class="rw-material-track">' + ''.join(material_cards) + '</div>', unsafe_allow_html=True)

    settings_col, action_col = st.columns([1.05, .95], gap="medium")
    with settings_col:
        with st.container(border=True):
            st.markdown('<div class="rw-panel-label">Prompt 设置</div><p class="rw-panel-copy">决定本次定制的表达方式，并选择本次允许调用的真实素材。</p>', unsafe_allow_html=True)
            length = st.selectbox("篇幅偏好", ["优先一页，必要时两页", "严格一页", "允许两页，保留细节"], key=f"rw_length_{job['id']}")
            default_materials = list(experience_options)[:min(6, len(experience_options))]
            selected_materials = st.multiselect("本次使用的 Experience Bank 素材", list(experience_options), default=default_materials, key=f"rw_materials_{job['id']}")
            notes = st.text_area("补充要求", placeholder="例如：突出 HR Tech、减少行政类描述……", height=92, key=f"rw_notes_{job['id']}")

    with action_col:
        with st.container(border=True):
            st.markdown('<div class="rw-panel-label">生成操作</div><p class="rw-panel-copy">生成 Prompt → 粘贴 AI 结果 → 解析预览 → 应用为右侧新版本。</p>', unsafe_allow_html=True)
            selected_ranked = [experience_options[label] for label in selected_materials]
            if st.button("生成定制 Prompt", key=f"rw_generate_{job['id']}", type="primary", use_container_width=True):
                prompt_settings = {"tone": "专业、清晰、精炼", "length": length, "notes": notes}
                prompt = _build_prompt(job, st.session_state[base_key], selected_ranked, prompt_settings)
                st.session_state[f"rw_prompt_{job['id']}"] = prompt
                save_resume_prompt(job["id"], None, prompt, [g["experience"].get("experience_code") for g in selected_ranked])
            prompt = st.session_state.get(f"rw_prompt_{job['id']}")
            if prompt:
                with st.expander("查看 Prompt", expanded=False):
                    st.text_area("Prompt", prompt, height=240, label_visibility="collapsed", key=f"rw_prompt_view_{job['id']}")
                    _copy_prompt_button(prompt)

            raw = st.text_area("AI 结果", key=f"rw_raw_{job['id']}", height=135, placeholder="粘贴 AI 生成的简历；普通文本或带导入标记的结果均可")
            if st.button("解析结果", key=f"rw_parse_{job['id']}", use_container_width=True):
                try:
                    parsed = _parse_import(raw)
                    st.session_state[f"rw_parsed_{job['id']}"] = parsed
                    save_import_record(None, raw, parsed, "parsed", "")
                except Exception as e:
                    save_import_record(None, raw, None, "failed", str(e))
                    st.error(str(e))

            parsed = st.session_state.get(f"rw_parsed_{job['id']}")
            if parsed:
                if parsed.get("_import_mode") == "plain_text":
                    st.markdown('<div class="rw-summary">已识别普通 AI 文本，可直接应用到 A4 编辑器。</div>', unsafe_allow_html=True)
                else:
                    sec = parsed.get("sections", {})
                    summary = "　".join(f"{k} {len(v) if isinstance(v,list) else 0}项" for k,v in sec.items())
                    st.markdown(f'<div class="rw-summary">{_safe(summary)}<br>待确认：{len(parsed.get("items_requiring_confirmation", []))} 项</div>', unsafe_allow_html=True)
                if st.button("应用到右侧并创建版本", key=f"rw_apply_{job['id']}", type="primary", use_container_width=True):
                    new_resume = _resume_from_import(parsed, st.session_state[base_key])
                    new_resume["_editor_settings"] = _editor_settings(st.session_state[generated_key])
                    new_resume["_editor_html"] = _resume_html(new_resume)
                    version_name = f'{job["company"]} {job["position"]} V{datetime.now():%m%d%H%M}'
                    version_id = create_resume_version(job["id"], version_name, new_resume, "ai_import", raw)
                    st.session_state[generated_key] = new_resume
                    save_import_record(version_id, raw, parsed, "applied", "")
                    st.session_state[f"rw_imported_{job['id']}"] = True
                    st.rerun()

            imported = st.session_state.get(f"rw_imported_{job['id']}", False) or bool(versions)
            if st.button("根据当前简历生成面试准备", key=f"rw_interview_{job['id']}", use_container_width=True, disabled=not imported):
                st.session_state[f"rw_pack_{job['id']}"] = _interview_pack(job, st.session_state[generated_key], selected_ranked or ranked)

    pack = st.session_state.get(f"rw_pack_{job['id']}")
    if pack:
        st.markdown('<div class="rw-interview"><h2>岗位专属 Interview Pack</h2>', unsafe_allow_html=True)
        tabs = st.tabs(["预测问题", "STAR 故事", "经历回链"])
        with tabs[0]:
            for q in pack["questions"]:
                st.markdown(f'<div class="rw-q"><b>{_safe(q["question"])}</b><br><small>来源：{_safe(q["experience"])} · {_safe(q["source"])}</small></div>', unsafe_allow_html=True)
        with tabs[1]:
            for story in pack["stories"]:
                st.markdown(f'''<div class="rw-star"><h4>{_safe(story['title'])} <small>({_safe(story['source'])})</small></h4><p><b>S</b> {_safe(story['s'],'')}</p><p><b>T</b> {_safe(story['t'],'')}</p><p><b>A</b> {_safe(story['a'],'')}</p><p><b>R</b> {_safe(story['r'],'')}</p><p><b>注意</b> {_safe(story['risk'],'无')}</p></div>''', unsafe_allow_html=True)
        with tabs[2]:
            for story in pack["stories"]:
                st.markdown(f'<div class="rw-q"><b>{_safe(story["experience"])}｜{_safe(story["title"])}</b><br>Experience Bank source_id：{_safe(story["source"])}</div>', unsafe_allow_html=True)
        st.markdown('</div>', unsafe_allow_html=True)


def _styles():
    st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Gaegu:wght@400;700&family=Noto+Sans+SC:wght@400;500;600;700&display=swap');
    :root{--resume-bg:#fef5f5;--resume-panel:#fffafa;--resume-panel-line:#f1e7e7}
    html,body,.stApp,[data-testid="stAppViewContainer"],[data-testid="stMain"]{min-height:100dvh!important;height:auto!important;background:var(--resume-bg)!important}
    [data-testid="stMain"]{padding:56px 14px 18px!important;overflow:visible!important;box-sizing:border-box!important}
    .block-container,[data-testid="stMainBlockContainer"]{display:block!important;width:100%!important;max-width:none!important;height:auto!important;max-height:none!important;min-height:calc(100dvh - 74px)!important;margin:0!important;padding:42px clamp(24px,3vw,54px) 38px!important;overflow:visible!important;border:1px solid var(--resume-panel-line)!important;border-radius:clamp(38px,4vw,66px)!important;background:var(--resume-panel)!important;box-shadow:0 20px 70px rgba(30,30,30,.055)!important;box-sizing:border-box!important}
    .block-container>[data-testid="stVerticalBlock"],[data-testid="stMainBlockContainer"]>[data-testid="stVerticalBlock"]{height:auto!important;max-height:none!important;overflow:visible!important}
    .rw-heading{max-width:900px;margin:-38px auto 20px;text-align:center}.rw-heading h1{margin:0;color:#111;font-family:'Gaegu','Noto Sans SC',sans-serif;font-size:clamp(3.2rem,5vw,5rem);font-weight:700;line-height:.93;letter-spacing:.04em}.rw-heading p{margin:13px 0 0;color:#858582;font:600 clamp(.78rem,1vw,.95rem)/1.55 'Noto Sans SC',sans-serif}.rw-heading p b{color:#444;font-weight:700}
    .rw-canvas-head{display:flex;align-items:center;justify-content:flex-end;margin:0 0 8px;padding:25px 2px 0}.rw-canvas-head span{font-size:10px;color:#727272;border:1px solid #dededb;border-radius:99px;padding:5px 9px;background:#fafaf9}.rw-output-note{font-size:10px;color:#777;text-align:center;margin:7px 0 0}
    .rw-sticky-anchor{display:none}div[data-testid="stColumn"]:has(.rw-sticky-anchor){position:sticky;top:4.3rem;align-self:flex-start;max-height:calc(100vh - 5.2rem);overflow-y:auto;margin-top:-18px;padding:12px 16px 16px!important;border:1px solid #efc49e;border-radius:22px;background:#ffe8d6;scrollbar-width:thin;scrollbar-color:#dc9b69 transparent;box-shadow:0 16px 38px rgba(126,76,36,.08)}div[data-testid="stColumn"]:has(.rw-sticky-anchor) div[data-testid="stVerticalBlockBorderWrapper"]{background:rgba(255,255,255,.62)!important}div[data-testid="stColumn"]:has(.rw-sticky-anchor) button[kind="primary"]{background:#dd7f46!important;border-color:#dd7f46!important;color:#fff!important}div[data-testid="stColumn"]:has(.rw-sticky-anchor) span[data-baseweb="tag"]{background:#f3b37f!important;color:#6f3c1d!important}div[data-testid="stColumn"]:has(.rw-sticky-anchor) span[data-baseweb="tag"] svg{fill:#6f3c1d!important;color:#6f3c1d!important}
    .rw-prompt-dock-anchor{display:none}div[data-testid="stLayoutWrapper"]:has(.rw-prompt-dock-anchor){position:fixed!important;left:auto!important;right:24px!important;width:min(350px,calc(100vw - 112px))!important;max-width:none!important;box-sizing:border-box!important;bottom:18px;z-index:99990!important;padding:12px 13px 13px!important;overflow:hidden;border:1px solid rgba(255,255,255,.86)!important;border-radius:12px!important;background:rgba(255,255,255,.62)!important;box-shadow:0 14px 32px rgba(50,48,38,.2),inset 0 1px 0 rgba(255,255,255,.76)!important;backdrop-filter:blur(17px) saturate(1.03);-webkit-backdrop-filter:blur(17px) saturate(1.03);transition:width .18s ease,box-shadow .18s ease;touch-action:none;font-family:'Kaiti SC','STKaiti','KaiTi','Gaegu',cursive!important}div[data-testid="stLayoutWrapper"]:has(.rw-prompt-dock-anchor)::before{content:'';position:absolute;z-index:0;left:0;right:0;top:0;height:57px;background:rgba(255,255,255,.3);border-bottom:1px solid rgba(255,255,255,.62);pointer-events:none}div[data-testid="stLayoutWrapper"]:has(.rw-prompt-dock-anchor)>div{position:relative;z-index:1}div[data-testid="stLayoutWrapper"]:has(.rw-prompt-dock-anchor):hover{box-shadow:0 17px 38px rgba(50,48,38,.24),inset 0 1px 0 rgba(255,255,255,.84)!important}div[data-testid="stLayoutWrapper"]:has(.rw-dock-collapsed){width:252px!important;padding:9px 10px!important;background:rgba(255,255,255,.62)!important}div[data-testid="stLayoutWrapper"]:has(.rw-dock-collapsed)::before{height:100%}div[data-testid="stLayoutWrapper"]:has(.rw-prompt-dock-anchor) .stTextArea{margin:1px 0 4px}div[data-testid="stLayoutWrapper"]:has(.rw-prompt-dock-anchor) .stTextArea textarea{min-height:96px!important;height:96px!important;resize:none!important;padding:7px 3px!important;background:transparent!important;border:0!important;border-bottom:1px dashed rgba(77,84,80,.3)!important;border-radius:0!important;box-shadow:none!important;line-height:1.65!important;font-family:'Kaiti SC','STKaiti','KaiTi',cursive!important;font-size:12px!important}div[data-testid="stLayoutWrapper"]:has(.rw-prompt-dock-anchor) button{min-height:36px!important;background:rgba(255,255,255,.28)!important;border-color:rgba(87,96,94,.16)!important;color:#343735!important;font-family:'Kaiti SC','STKaiti','KaiTi',cursive!important}div[data-testid="stLayoutWrapper"]:has(.rw-dock-expanded) button{background:transparent!important;border:0!important;border-radius:0!important;border-top:1px solid rgba(77,84,80,.18)!important;box-shadow:none!important}div[data-testid="stLayoutWrapper"]:has(.rw-dock-expanded) button[kind="primary"]{background:transparent!important;border-color:rgba(77,84,80,.18)!important;color:#263638!important;font-weight:750!important}.rw-dock-title{padding:2px 0 5px;cursor:grab;user-select:none;-webkit-user-select:none}.rw-dock-title:active{cursor:grabbing}.rw-dock-title b{display:block;font-size:15px;color:#252a28;font-family:'Kaiti SC','STKaiti','KaiTi','Gaegu',cursive}.rw-dock-title b::before{content:'⠿';display:inline-block;margin-right:6px;color:rgba(71,82,79,.48);font-size:13px}.rw-dock-title span{display:block;font-size:10px;color:#676e6b;margin-top:3px;line-height:1.35}.rw-note-step{display:grid;grid-template-columns:19px 1fr;column-gap:7px;align-items:center;margin:8px 0 2px;color:#333833;font-family:'Kaiti SC','STKaiti','KaiTi',cursive}.rw-note-step i{display:grid;place-items:center;width:18px;height:18px;border:1px solid rgba(69,78,74,.28);border-radius:50%;font:700 11px/1 'Gaegu',cursive;font-style:normal}.rw-note-step b{font-size:12px}.rw-note-step span{grid-column:2;font-size:9px;color:#7a7e79;margin-top:1px}.rw-selected-quote{margin:3px 0 4px;padding:3px 2px 5px 26px;border:0;border-bottom:1px dashed rgba(77,84,80,.22);background:transparent;color:#505653;font-family:'Kaiti SC','STKaiti','KaiTi',cursive;font-size:11px;line-height:1.5;max-height:54px;overflow:auto}.rw-note-empty{color:#898b87}.block-container{padding-bottom:6rem!important}
    .rw-panel-intro{padding:0 2px 3px}.rw-panel-intro b{display:block;font-family:'Gaegu','Noto Sans SC',sans-serif;font-size:17px;line-height:1;font-weight:700;letter-spacing:.05em;text-transform:lowercase;color:#7c573e}.rw-change-stat{display:grid;grid-template-columns:1fr 1fr;gap:7px;margin:8px 0 11px}.rw-change-stat div:first-child{background:#d9f1e1}.rw-change-stat div:last-child{background:#dceefa}.rw-change-stat div{border-radius:11px;padding:9px}.rw-change-stat strong{display:block;font-size:18px}.rw-change-stat small{font-size:9px;color:#66706a}.rw-change-item{font-size:10px;line-height:1.45;padding:8px 9px;border:1px solid rgba(30,30,30,.06);border-radius:9px;background:rgba(255,255,255,.8);margin:5px 0}.rw-change-item.added{border-left:3px solid #68a884}.rw-change-item.removed{border-left:3px solid #7ba8d8;color:#5f6870}.rw-comment-card{padding:9px;border:1px solid #efc49e;border-radius:11px;background:rgba(255,255,255,.78);margin:7px 0}.rw-comment-card b{display:block;font-size:9px;color:#a46135;margin-bottom:4px}.rw-comment-card p{font-size:10px;line-height:1.45;margin:2px 0;color:#4e443d}
    .rw-material-mini{border:1px solid #dededb;border-radius:12px;background:#fff;padding:10px;margin:7px 0}.rw-material-mini header{display:flex;justify-content:space-between;gap:7px;font-size:9px;color:#777}.rw-material-mini b{display:block;font-size:11px;margin:6px 0 2px}.rw-material-mini p{font-size:10px;color:#777;line-height:1.45;margin:0}.rw-summary{padding:10px 11px;border:1px solid #dededb;border-radius:11px;background:#f5f5f3;font-size:10px;line-height:1.6;margin:8px 0}
    .rw-compare-head{display:flex;justify-content:space-between;align-items:center;margin:8px 0}.rw-compare-head b{font-size:13px}.rw-compare-head span{font-size:9px;color:#777;border-radius:99px;background:#efefed;padding:4px 8px}.rw-interview{margin-top:18px;border:1px solid #dededb;border-radius:18px;padding:18px;background:#f7f7f5}.rw-q{background:#fff;border:1px solid #e6e6e3;padding:12px;border-radius:11px;margin:8px 0;font-size:12px}.rw-star{background:#fff;border:1px solid #e6e6e3;padding:15px;border-radius:12px;margin:10px 0}.rw-star p{font-size:11px;line-height:1.65;margin:5px 0}
    div[data-testid="stCustomComponentV1"]{border-radius:18px;overflow:hidden}div[data-testid="stVerticalBlockBorderWrapper"]{border-color:#dededb!important;border-radius:16px!important;background:#fafaf9!important}div[data-testid="stTabs"] button[role="tab"]{font-size:11px!important;padding-left:9px!important;padding-right:9px!important}
    @media(max-width:1100px){[data-testid="stMain"]{padding:48px 8px 12px!important}.block-container,[data-testid="stMainBlockContainer"]{min-height:calc(100dvh - 60px)!important;padding:46px 16px 26px!important;border-radius:34px!important}div[data-testid="stColumn"]:has(.rw-sticky-anchor){position:static;max-height:none;overflow:visible;margin-top:0}.rw-heading{align-items:flex-start;flex-direction:column;margin:-20px auto 20px}div[data-testid="stLayoutWrapper"]:has(.rw-prompt-dock-anchor){right:14px!important;width:min(350px,calc(100vw - 100px))!important}div[data-testid="stLayoutWrapper"]:has(.rw-dock-collapsed){width:240px!important}}
    /* Stable resume studio: contain both columns and keep every control in-flow. */
    html,body,.stApp,[data-testid="stAppViewContainer"],[data-testid="stMain"]{overflow-x:hidden!important}
    [data-testid="stMain"],.block-container,[data-testid="stMainBlockContainer"]{overflow-x:hidden!important}
    [data-testid="stHorizontalBlock"],[data-testid="stLayoutWrapper"],div[data-testid="stColumn"]{min-width:0!important;max-width:100%!important;box-sizing:border-box!important}
    div[data-testid="stColumn"]:has(.rw-sticky-anchor){min-width:0!important;max-width:390px!important;overflow-x:hidden!important}
    div[data-testid="stCustomComponentV1"]{display:block!important;width:100%!important;max-width:100%!important;min-width:0!important}
    div[data-testid="stCustomComponentV1"] iframe{display:block!important;width:100%!important;max-width:100%!important;min-height:1040px!important;border:0!important}
    .rw-prompt-dock-anchor,div[data-testid="stLayoutWrapper"]:has(.rw-prompt-dock-anchor){display:none!important}
    .block-container{padding-bottom:3rem!important}
    </style>
    """, unsafe_allow_html=True)


def _enable_note_drag(job_id):
    """Attach bounded drag behavior to the floating note without changing app data."""
    storage_key = json.dumps(f"careerpilot_resume_note_position_{job_id}")
    components.html(
        f"""
        <script>
        (() => {{
          const host = window.parent;
          const doc = host.document;
          const stateKey = '__careerPilotResumeNoteDrag';
          const storageKey = {storage_key};
          let attempts = 0;

          function init() {{
            const marker = doc.querySelector('.rw-prompt-dock-anchor');
            const dock = marker && marker.closest('[data-testid="stLayoutWrapper"]');
            const handle = dock && dock.querySelector('.rw-dock-title');
            if (!dock || !handle) {{
              if (attempts++ < 40) host.setTimeout(init, 120);
              return;
            }}

            if (host[stateKey] && host[stateKey].cleanup) host[stateKey].cleanup();
            let drag = null;

            const clamp = (value, low, high) => Math.min(Math.max(value, low), Math.max(low, high));
            const setPosition = (x, y) => {{
              const minX = host.innerWidth < 760 ? 12 : 92;
              const maxX = host.innerWidth - dock.getBoundingClientRect().width - 12;
              const maxY = host.innerHeight - dock.getBoundingClientRect().height - 12;
              const nextX = clamp(Number(x) || minX, minX, maxX);
              const nextY = clamp(Number(y) || 12, 12, maxY);
              dock.style.setProperty('left', `${{nextX}}px`, 'important');
              dock.style.setProperty('top', `${{nextY}}px`, 'important');
              dock.style.setProperty('right', 'auto', 'important');
              dock.style.setProperty('bottom', 'auto', 'important');
              return {{x: nextX, y: nextY}};
            }};

            const savePosition = position => {{
              try {{ host.localStorage.setItem(storageKey, JSON.stringify(position)); }} catch (_) {{}}
            }};
            const restorePosition = () => {{
              try {{
                const saved = JSON.parse(host.localStorage.getItem(storageKey) || 'null');
                if (saved && Number.isFinite(saved.x) && Number.isFinite(saved.y)) setPosition(saved.x, saved.y);
              }} catch (_) {{}}
            }};

            const onPointerDown = event => {{
              if (event.button !== 0) return;
              const rect = dock.getBoundingClientRect();
              drag = {{
                pointerId: event.pointerId,
                startX: event.clientX,
                startY: event.clientY,
                left: rect.left,
                top: rect.top,
              }};
              dock.style.transition = 'none';
              dock.style.willChange = 'left, top, transform';
              dock.style.transform = 'scale(1.02) rotate(-0.35deg)';
              handle.setPointerCapture?.(event.pointerId);
              event.preventDefault();
            }};
            const onPointerMove = event => {{
              if (!drag || event.pointerId !== drag.pointerId) return;
              setPosition(
                drag.left + event.clientX - drag.startX,
                drag.top + event.clientY - drag.startY,
              );
              const tilt = clamp((event.clientX - drag.startX) / 110, -1.1, 1.1);
              dock.style.transform = `scale(1.02) rotate(${{tilt}}deg)`;
              event.preventDefault();
            }};
            const onPointerUp = event => {{
              if (!drag || event.pointerId !== drag.pointerId) return;
              const rect = dock.getBoundingClientRect();
              const settledLeft = Number.parseFloat(dock.style.left);
              const settledTop = Number.parseFloat(dock.style.top);
              savePosition(setPosition(
                Number.isFinite(settledLeft) ? settledLeft : rect.left,
                Number.isFinite(settledTop) ? settledTop : rect.top,
              ));
              dock.style.transition = 'transform .16s ease, box-shadow .16s ease';
              dock.style.transform = '';
              host.setTimeout(() => {{
                dock.style.transition = '';
                dock.style.willChange = '';
              }}, 170);
              drag = null;
            }};
            const onResize = () => {{
              const rect = dock.getBoundingClientRect();
              savePosition(setPosition(rect.left, rect.top));
            }};

            handle.addEventListener('pointerdown', onPointerDown);
            doc.addEventListener('pointermove', onPointerMove, true);
            doc.addEventListener('pointerup', onPointerUp, true);
            doc.addEventListener('pointercancel', onPointerUp, true);
            host.addEventListener('resize', onResize);
            restorePosition();
            host.setTimeout(restorePosition, 180);

            host[stateKey] = {{
              cleanup: () => {{
                handle.removeEventListener('pointerdown', onPointerDown);
                doc.removeEventListener('pointermove', onPointerMove, true);
                doc.removeEventListener('pointerup', onPointerUp, true);
                doc.removeEventListener('pointercancel', onPointerUp, true);
                host.removeEventListener('resize', onResize);
              }}
            }};
          }}

          init();
        }})();
        </script>
        """,
        height=0,
        width=0,
    )


def render_resume_workspace(jobs):
    _styles()
    payload = _experience_payload()
    if not jobs:
        st.warning("请先在“职位申请”中创建目标岗位。")
        return

    st.markdown('''<div class="rw-heading"><h1>resume tailoring</h1></div>''', unsafe_allow_html=True)
    top_left, top_right = st.columns([1.35, 1], gap="medium")
    job_labels = {f'{j["company"]}｜{j["position"]}': j for j in jobs}
    selected_label = top_left.selectbox("选择目标岗位 / JD", list(job_labels), key="rw_job_v2")
    view_mode = top_right.radio("查看模式", ["编辑", "审阅", "干净预览", "全屏对比"], horizontal=True, key="rw_view_mode", label_visibility="visible")
    job = job_labels[selected_label]
    ranked = _rank_experiences(job, payload)
    versions = get_resume_versions(job["id"])
    base_version = next((v for v in versions if v.get("version_type") == "base_edit"), None)
    generated_version = next((v for v in versions if v.get("version_type") != "base_edit"), None)

    base_key = f'rw_base_{job["id"]}'
    generated_key = f'rw_generated_{job["id"]}'
    editor_revision_key = f"rw_editor_revision_{job['id']}"
    if editor_revision_key not in st.session_state:
        st.session_state[editor_revision_key] = 0
    if base_key not in st.session_state:
        if base_version:
            st.session_state[base_key] = json.loads(base_version["content_json"])
        elif PERSONAL_RESUME_BASE:
            st.session_state[base_key] = deepcopy(PERSONAL_RESUME_BASE)
        else:
            st.session_state[base_key] = _default_resume(payload)
    if generated_key not in st.session_state:
        st.session_state[generated_key] = json.loads(generated_version["content_json"]) if generated_version else deepcopy(st.session_state[base_key])
    demo_version = (PERSONAL_RESUME_BASE or {}).get("_demo_version")
    if demo_version and (
        st.session_state[base_key].get("_demo_version") != demo_version
        or st.session_state[generated_key].get("_demo_version") != demo_version
    ):
        st.session_state[base_key] = deepcopy(PERSONAL_RESUME_BASE)
        st.session_state[generated_key] = deepcopy(PERSONAL_RESUME_BASE)
        st.session_state[editor_revision_key] += 1
    base_resume = st.session_state[base_key]
    generated_resume = st.session_state[generated_key]
    template_resume = base_resume if base_resume.get("sections") else (deepcopy(PERSONAL_RESUME_BASE) if PERSONAL_RESUME_BASE else None)
    current_editor_html = _editor_value(generated_resume)
    needs_template_repair = (
        not generated_resume.get("sections")
        or "个人概述" in current_editor_html
        or "mailto:" in current_editor_html
    )
    if template_resume and needs_template_repair:
        migrated = _resume_from_import(
            _legacy_editor_import_data(current_editor_html),
            template_resume,
        )
        migrated["_editor_settings"] = _editor_settings(generated_resume)
        migrated["_review_comments"] = deepcopy(generated_resume.get("_review_comments") or [])
        migrated["_editor_html"] = _resume_html(migrated)
        st.session_state[generated_key] = migrated
        generated_resume = migrated
        st.session_state[editor_revision_key] += 1
    selection_key = f"rw_selected_text_{job['id']}"
    comments_key = f"rw_review_comments_{job['id']}"
    comment_draft_key = f"rw_comment_draft_{job['id']}"
    dock_open_key = f"rw_note_dock_open_{job['id']}"
    handled_selection_key = f"rw_handled_selection_{job['id']}"
    applied_notice_key = f"rw_applied_notice_{job['id']}"
    if selection_key not in st.session_state:
        st.session_state[selection_key] = ""
    if comments_key not in st.session_state:
        st.session_state[comments_key] = deepcopy(generated_resume.get("_review_comments") or [])
    if comment_draft_key not in st.session_state:
        st.session_state[comment_draft_key] = ""
    if dock_open_key not in st.session_state:
        st.session_state[dock_open_key] = False
    if handled_selection_key not in st.session_state:
        st.session_state[handled_selection_key] = 0
    if st.session_state.pop(applied_notice_key, False):
        st.toast("已应用到 A4 编辑器，并创建新的岗位版本。")
    base_html, generated_html = _editor_value(base_resume), _editor_value(generated_resume)
    changes = _change_summary(base_html, generated_html)
    file_name_key = f"rw_file_name_{job['id']}"
    if file_name_key not in st.session_state:
        default_file_name = f'{job.get("company", "")}_{job.get("position", "")}_定制简历'
        st.session_state[file_name_key] = (
            generated_version.get("version_name") if generated_version else ""
        ) or default_file_name

    material_options = {}
    for group in ranked:
        exp = group["experience"]
        code = exp.get("experience_code") or f'exp_{exp["id"]}'
        label = f'{exp.get("organization") or "未命名经历"}｜{exp.get("role") or "经历"} [{code}]'
        material_options[label] = group
    material_key = f"rw_materials_v2_{job['id']}"
    if material_key not in st.session_state:
        st.session_state[material_key] = list(material_options)[:min(6, len(material_options))]
    dock_note_key = f"rw_dock_note_{job['id']}"
    if dock_note_key not in st.session_state:
        st.session_state[dock_note_key] = st.session_state.get(f"rw_notes_v2_{job['id']}", "")

    def capture_editor_selection(result):
        if isinstance(result, dict):
            selected = str(result.get("selected_text") or "").strip()
            selection_id = int(result.get("selection_id") or 0)
            if selected and selection_id and selection_id != st.session_state.get(handled_selection_key):
                st.session_state[handled_selection_key] = selection_id
                st.session_state[selection_key] = selected
                st.session_state[dock_open_key] = True

    def sync_comments_to_resume():
        current = deepcopy(st.session_state[generated_key])
        current["_review_comments"] = deepcopy(st.session_state.get(comments_key, []))
        st.session_state[generated_key] = current

    def add_review_comment():
        quote = str(st.session_state.get(selection_key) or "").strip()
        note = str(st.session_state.get(comment_draft_key) or "").strip()
        if not quote or not note:
            return
        st.session_state[comments_key].append({
            "quote": quote,
            "note": note,
            "created_at": datetime.now().isoformat(timespec="minutes"),
        })
        st.session_state[comment_draft_key] = ""
        st.session_state[selection_key] = ""
        st.session_state[dock_open_key] = False
        sync_comments_to_resume()

    def open_note_dock():
        st.session_state[dock_open_key] = True

    def close_note_dock():
        st.session_state[dock_open_key] = False

    if view_mode == "全屏对比":
        st.caption("全屏对比仅用于审阅，不影响最终 A4 与导出文件。")
        compare_left, compare_right = st.columns(2, gap="medium")
        with compare_left:
            st.markdown('<div class="rw-compare-head"><b>基础简历</b><span>只读</span></div>', unsafe_allow_html=True)
            base_compare_result = resume_editor(base_html, _editor_settings(base_resume), key=f"rw_compare_base_{job['id']}", height=980, mode="compare")
            capture_editor_selection(base_compare_result)
        with compare_right:
            st.markdown('<div class="rw-compare-head"><b>当前岗位版本</b><span>只读</span></div>', unsafe_allow_html=True)
            generated_compare_result = resume_editor(generated_html, _editor_settings(generated_resume), key=f"rw_compare_generated_{job['id']}", height=980, mode="review", base_value=base_html)
            capture_editor_selection(generated_compare_result)
    else:
        main_col, side_col = st.columns([1, .34], gap="medium")
        with main_col:
            mode_map = {"编辑": "edit", "审阅": "review", "干净预览": "preview"}
            file_title_col, canvas_status_col = st.columns([1, .3], vertical_alignment="bottom")
            export_name = file_title_col.text_input(
                "文件标题",
                key=file_name_key,
                help="同时用于保存当前岗位版本以及 PDF、Word 的导出文件名。",
                placeholder="例如：Northstar_PeopleOps_定制简历",
            ).strip()
            canvas_status_col.markdown(f'<div class="rw-canvas-head"><span>{_safe(view_mode)} · A4 210 × 297 mm</span></div>', unsafe_allow_html=True)
            export_name = re.sub(r'\.(?:pdf|docx)$', '', export_name, flags=re.IGNORECASE).strip()
            export_name = export_name or f'{job.get("company", "")}_{job.get("position", "")}_定制简历'
            clean_name = re.sub(r'[^\w\u4e00-\u9fff-]+', '_', export_name).strip('_') or "定制简历"
            editor_result = resume_editor(
                generated_html,
                _editor_settings(generated_resume),
                key=f"rw_main_editor_{job['id']}_{mode_map[view_mode]}_{st.session_state[editor_revision_key]}",
                height=1040,
                mode=mode_map[view_mode],
                base_value=base_html,
                document_version=f"{demo_version or 'resume'}:{st.session_state[editor_revision_key]}",
            )
            capture_editor_selection(editor_result)
            if view_mode == "编辑":
                st.session_state[generated_key] = _apply_editor_result(generated_resume, editor_result)
                generated_resume = st.session_state[generated_key]
            st.markdown('<div class="rw-output-note">PDF 打印会自动隐藏审阅标记；Word 与 PDF 均使用 A4 纸张设置。</div>', unsafe_allow_html=True)
            save_col, pdf_col, word_col = st.columns(3)
            if save_col.button("保存当前岗位版本", key=f"rw_save_v2_{job['id']}", use_container_width=True):
                create_resume_version(job["id"], export_name, st.session_state[generated_key], "manual")
                st.success("岗位版本及排版设置已保存。")
            pdf_col.download_button(
                "导出 PDF (.pdf)",
                data=_pdf_bytes(st.session_state[generated_key]),
                file_name=f"{clean_name}.pdf",
                mime="application/pdf",
                key=f"rw_pdf_{job['id']}",
                use_container_width=True,
            )
            word_col.download_button(
                "导出 Word (.docx)",
                data=_docx_bytes(st.session_state[generated_key]),
                file_name=f"{clean_name}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                key=f"rw_word_{job['id']}",
                use_container_width=True,
            )

        with side_col:
            st.markdown('<span class="rw-sticky-anchor"></span><div class="rw-panel-intro"><b>tailoring desk</b></div>', unsafe_allow_html=True)
            prompt_tab, changes_tab, materials_tab = st.tabs(["Prompt", "修改", "素材"])
            with materials_tab:
                st.multiselect("本次使用的 Experience Bank 素材", list(material_options), key=material_key)
                for label in st.session_state.get(material_key, [])[:6]:
                    group = material_options.get(label)
                    if not group:
                        continue
                    exp = group["experience"]
                    tags = " · ".join(group.get("matches") or []) or "等待关键词匹配"
                    st.markdown(f'''<div class="rw-material-mini"><header><span>{_safe(exp.get("experience_code"), "EXP")}</span><span>匹配度 {group.get("score",0)}%</span></header><b>{_safe(exp.get("organization"))}</b><p>{_safe(exp.get("role"))}<br>{_safe(group.get("match_reason"))}<br>{_safe(tags)}</p></div>''', unsafe_allow_html=True)

            with changes_tab:
                st.markdown(f'<div class="rw-change-stat"><div><strong>{len(changes["added"])}</strong><small>新增 / 改写</small></div><div><strong>{len(changes["removed"])}</strong><small>删除 / 压缩</small></div></div>', unsafe_allow_html=True)
                for item in changes["added"][:7]:
                    st.markdown(f'<div class="rw-change-item added">{_safe(item)}</div>', unsafe_allow_html=True)
                for item in changes["removed"][:5]:
                    st.markdown(f'<div class="rw-change-item removed">{_safe(item)}</div>', unsafe_allow_html=True)
                comments = st.session_state.get(comments_key, [])
                if comments:
                    st.caption(f"批注 · {len(comments)} 条（仅进入 Prompt，不进入导出）")
                for index, comment in enumerate(comments):
                    st.markdown(f'<div class="rw-comment-card"><b>原文</b><p>{_safe(comment.get("quote"), "")}</p><b>修改要求</b><p>{_safe(comment.get("note"), "")}</p></div>', unsafe_allow_html=True)
                    if st.button("删除这条批注", key=f"rw_delete_comment_{job['id']}_{index}", use_container_width=True):
                        st.session_state[comments_key].pop(index)
                        sync_comments_to_resume()
                        st.rerun()
                if st.button("恢复为基础简历", key=f"rw_restore_base_{job['id']}", use_container_width=True):
                    restored = deepcopy(st.session_state[base_key])
                    restored["_editor_settings"] = _editor_settings(st.session_state[generated_key])
                    restored["_review_comments"] = []
                    st.session_state[generated_key] = restored
                    st.session_state[comments_key] = []
                    st.session_state[selection_key] = ""
                    st.session_state[editor_revision_key] += 1
                    st.rerun()
                if generated_version and st.button("恢复最近保存版本", key=f"rw_restore_saved_{job['id']}", use_container_width=True):
                    restored = json.loads(generated_version["content_json"])
                    st.session_state[generated_key] = restored
                    st.session_state[comments_key] = deepcopy(restored.get("_review_comments") or [])
                    st.session_state[selection_key] = ""
                    st.session_state[editor_revision_key] += 1
                    st.rerun()

            with prompt_tab:
                length = st.selectbox("篇幅偏好", ["优先一页，必要时两页", "严格一页", "允许两页，保留细节"], key=f"rw_length_v2_{job['id']}")
                notes = st.text_area(
                    "补充要求",
                    key=dock_note_key,
                    height=92,
                    placeholder="例如：突出业务影响，减少流程描述……",
                )
                selected_ranked = [material_options[label] for label in st.session_state.get(material_key, []) if label in material_options]
                if st.button("生成定制 Prompt", key=f"rw_generate_v2_{job['id']}", type="primary", use_container_width=True):
                    prompt_settings = {"tone": "专业、清晰、精炼", "length": length, "notes": _prompt_notes(notes, st.session_state.get(comments_key, []))}
                    prompt = _build_prompt(job, st.session_state[base_key], selected_ranked, prompt_settings)
                    st.session_state[f"rw_prompt_{job['id']}"] = prompt
                    save_resume_prompt(job["id"], None, prompt, [g["experience"].get("experience_code") for g in selected_ranked])
                prompt = st.session_state.get(f"rw_prompt_{job['id']}")
                if prompt:
                    with st.expander("查看 Prompt", expanded=False):
                        st.text_area("Prompt", prompt, height=220, label_visibility="collapsed", key=f"rw_prompt_view_v2_{job['id']}")
                        _copy_prompt_button(prompt)
                raw = st.text_area("AI 结果", key=f"rw_raw_v2_{job['id']}", height=125, placeholder="粘贴 AI 生成的简历；普通文本或带导入标记的结果均可")
                if st.button("解析结果", key=f"rw_parse_v2_{job['id']}", use_container_width=True):
                    try:
                        parsed = _parse_import(raw)
                        st.session_state[f"rw_parsed_{job['id']}"] = parsed
                        save_import_record(None, raw, parsed, "parsed", "")
                    except Exception as exc:
                        save_import_record(None, raw, None, "failed", str(exc))
                        st.error(str(exc))
                parsed = st.session_state.get(f"rw_parsed_{job['id']}")
                if parsed:
                    if parsed.get("_import_mode") == "plain_text":
                        st.caption("已识别普通 AI 文本，可直接应用到 A4 编辑器。")
                    else:
                        sec = parsed.get("sections", {})
                        summary = "　".join(f"{key} {len(value) if isinstance(value,list) else 0}项" for key,value in sec.items())
                        st.markdown(f'<div class="rw-summary">{_safe(summary)}<br>待确认：{len(parsed.get("items_requiring_confirmation", []))} 项</div>', unsafe_allow_html=True)
                    if st.button("应用为新的岗位版本", key=f"rw_apply_v2_{job['id']}", type="primary", use_container_width=True):
                        new_resume = _resume_from_import(parsed, st.session_state[base_key])
                        new_resume["_editor_settings"] = _editor_settings(st.session_state[generated_key])
                        new_resume["_editor_html"] = _resume_html(new_resume)
                        new_resume["_review_comments"] = []
                        version_name = f'{job["company"]} {job["position"]} V{datetime.now():%m%d%H%M}'
                        version_id = create_resume_version(job["id"], version_name, new_resume, "ai_import", raw)
                        st.session_state[generated_key] = new_resume
                        st.session_state[comments_key] = []
                        st.session_state[selection_key] = ""
                        st.session_state[editor_revision_key] += 1
                        st.session_state[applied_notice_key] = True
                        save_import_record(version_id, raw, parsed, "applied", "")
                        st.session_state[f"rw_imported_{job['id']}"] = True
                        st.rerun()
                imported = st.session_state.get(f"rw_imported_{job['id']}", False) or bool(versions)
                if st.button("生成面试准备", key=f"rw_interview_v2_{job['id']}", use_container_width=True, disabled=not imported):
                    st.session_state[f"rw_pack_{job['id']}"] = _interview_pack(job, st.session_state[generated_key], selected_ranked or ranked)

    dock_ranked = [material_options[label] for label in st.session_state.get(material_key, []) if label in material_options]
    selected_quote = str(st.session_state.get(selection_key) or "").strip()
    active_note_key = comment_draft_key if selected_quote else dock_note_key
    dock_is_open = bool(st.session_state.get(dock_open_key))
    with st.container(border=True):
        if not dock_is_open:
            st.markdown('<div class="rw-prompt-dock-anchor rw-dock-collapsed">&nbsp;</div>', unsafe_allow_html=True)
            collapsed_title, collapsed_action = st.columns([1, .34], gap="small")
            badge = f" · 已选中" if selected_quote else ""
            collapsed_title.markdown(f'<div class="rw-dock-title"><b>👩🏻‍💼 简历批注{badge}</b><span>框选原文，随手加入 Prompt</span></div>', unsafe_allow_html=True)
            collapsed_action.button("＋", key=f"rw_open_dock_{job['id']}", use_container_width=True, on_click=open_note_dock)
        else:
            st.markdown('<div class="rw-prompt-dock-anchor rw-dock-expanded">&nbsp;</div>', unsafe_allow_html=True)
            dock_title, dock_close_action = st.columns([1, .22], gap="small")
            if selected_quote:
                quote_preview = selected_quote if len(selected_quote) <= 80 else selected_quote[:80] + "…"
                dock_title.markdown('<div class="rw-dock-title"><b>👩🏻‍💼 简历批注</b><span>已选中原文，写下这段的修改要求</span></div>', unsafe_allow_html=True)
            else:
                dock_title.markdown('<div class="rw-dock-title"><b>👩🏻‍💼 简历批注</b><span>拖选简历文字即可添加逐段批注</span></div>', unsafe_allow_html=True)
            dock_close_action.button("×", key=f"rw_close_dock_{job['id']}", use_container_width=True, on_click=close_note_dock)
            st.markdown('<div class="rw-note-step"><i>1</i><b>框选原文</b><span>在 A4 简历中拖选需要调整的句子</span></div>', unsafe_allow_html=True)
            if selected_quote:
                st.markdown(f'<div class="rw-selected-quote">“{_safe(quote_preview)}”</div>', unsafe_allow_html=True)
            else:
                st.markdown('<div class="rw-selected-quote rw-note-empty">尚未框选原文，也可以直接记录整体要求。</div>', unsafe_allow_html=True)
            st.markdown('<div class="rw-note-step"><i>2</i><b>填写备注</b><span>写下希望 AI 如何修改</span></div>', unsafe_allow_html=True)
            dock_note = st.text_area(
                "批注或 Prompt 备注",
                key=active_note_key,
                placeholder="例如：突出业务影响，减少流程描述……" if selected_quote else "填写整体要求，或先拖选上方原文再逐段批注。",
                height=112,
                label_visibility="collapsed",
            )
            st.markdown('<div class="rw-note-step"><i>3</i><b>加入 Prompt</b><span>批注仅用于生成，不会进入导出文件</span></div>', unsafe_allow_html=True)
            dock_comment_action, dock_generate_action = st.columns(2, gap="small")
            dock_comment_action.button(
                "加入 Prompt",
                key=f"rw_add_comment_{job['id']}",
                use_container_width=True,
                disabled=not (selected_quote and str(dock_note or "").strip()),
                on_click=add_review_comment,
            )
            if dock_generate_action.button("生成 Prompt", key=f"rw_dock_generate_{job['id']}", type="primary", use_container_width=True):
                prompt_settings = {
                    "tone": "专业、清晰、精炼",
                    "length": st.session_state.get(f"rw_length_v2_{job['id']}", "优先一页，必要时两页"),
                    "notes": _prompt_notes(st.session_state.get(dock_note_key, ""), st.session_state.get(comments_key, [])),
                }
                prompt = _build_prompt(job, st.session_state[base_key], dock_ranked, prompt_settings)
                st.session_state[f"rw_prompt_{job['id']}"] = prompt
                save_resume_prompt(job["id"], None, prompt, [group["experience"].get("experience_code") for group in dock_ranked])

    # Floating prompt notes were removed to keep the two-column canvas stable.

    pack = st.session_state.get(f"rw_pack_{job['id']}")
    if pack:
        st.markdown('<div class="rw-interview"><h2>岗位专属 Interview Pack</h2>', unsafe_allow_html=True)
        tabs = st.tabs(["预测问题", "STAR 故事", "经历回链"])
        with tabs[0]:
            for question in pack["questions"]:
                st.markdown(f'<div class="rw-q"><b>{_safe(question["question"])}</b><br><small>来源：{_safe(question["experience"])} · {_safe(question["source"])}</small></div>', unsafe_allow_html=True)
        with tabs[1]:
            for story in pack["stories"]:
                st.markdown(f'''<div class="rw-star"><h4>{_safe(story['title'])}</h4><p><b>S</b> {_safe(story['s'],'')}</p><p><b>T</b> {_safe(story['t'],'')}</p><p><b>A</b> {_safe(story['a'],'')}</p><p><b>R</b> {_safe(story['r'],'')}</p></div>''', unsafe_allow_html=True)
        with tabs[2]:
            for story in pack["stories"]:
                st.markdown(f'<div class="rw-q"><b>{_safe(story["experience"])}｜{_safe(story["title"])}</b><br>source_id：{_safe(story["source"])}</div>', unsafe_allow_html=True)
        st.markdown('</div>', unsafe_allow_html=True)
